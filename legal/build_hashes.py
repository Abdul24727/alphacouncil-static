"""Generate canonical content hashes from the NORMALIZED authoritative legal source (.docx),
not from volatile HTML. Writes backend/legal_hashes.json (committed) — the acceptance records
and the build check read from it. build_id is content-derived (deterministic, no timestamp).
"""
import os, json, hashlib
from docx import Document

SRC = r"C:\projects\alphacouncil-static\legal\source"
OUT = r"C:\projects\alphacouncil\Argaam-council-before-testing-cluade-project-\llm-council\backend\legal_hashes.json"

# key -> source filename (matches legal_documents.LEGAL_DOCUMENTS)
DOCS = {
    "terms": "AlphaCouncil-Terms-of-Service-FINAL.docx",
    "privacy": "AlphaCouncil-Privacy-Policy-FINAL.docx",
    "refunds": "AlphaCouncil-Refund-Cancellation-Policy-FINAL.docx",
    "cookies": "AlphaCouncil-Cookie-Policy-FINAL.docx",
    "ai-financial-disclaimer": "AlphaCouncil-AI-Financial-Disclaimer-FINAL.docx",
}
COMBINED = "AlphaCouncil-Legal-Suite-FINAL.docx"

# Acceptance sentences (must match legal_documents.ACCEPTANCE_TEXT)
ACCEPTANCE_TEXT = {
    "registration": "I agree to the Terms of Service and acknowledge the Privacy Policy.",
    "checkout": ("I agree to the Terms of Service and acknowledge the Privacy Policy, "
                 "Refund and Cancellation Policy, and AI and Financial Disclaimer."),
}
# Document versions this hash-set is bound to (change the hash -> must change the version).
# v1.1 (2026-07-23): non-renewing-launch amendments to terms/privacy/refunds (Terms 10.2,
# Refunds 5, Privacy payment-services bullet) — DRAFT pending counsel approval/renumbering.
VERSIONS = {"terms": "1.1", "privacy": "1.1", "refunds": "1.1",
            "cookies": "1.0", "ai-financial-disclaimer": "1.0"}


# Presentation-only styles excluded from the CANONICAL content (they are rendered outside the
# substantive #legal-body: the <h1> title and the version/effective-date metadata line).
_PRESENTATION = {"Title", "Legal Metadata"}


def canonical_blocks(path):
    """Ordered list of substantive content blocks (headings, paragraphs, definitions, list
    items, then table cells) — the canonical artifact the rendered #legal-body must reproduce."""
    d = Document(path)
    blocks = []
    for p in d.paragraphs:
        st = (p.style.name if p.style else "") or ""
        if st in _PRESENTATION:
            continue
        if p.text.strip():
            blocks.append(p.text.strip())
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                if c.text.strip():
                    blocks.append(c.text.strip())
    return blocks


def normalized_text(path):
    return "\n".join(canonical_blocks(path))


def sha(text):
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


doc_hashes = {k: sha(normalized_text(os.path.join(SRC, f))) for k, f in DOCS.items()}
combined_hash = sha(normalized_text(os.path.join(SRC, COMBINED)))
accept_hashes = {k: sha(v) for k, v in ACCEPTANCE_TEXT.items()}

# Deterministic build id = hash of the ordered hash set + versions (no wall-clock).
build_material = json.dumps({"d": doc_hashes, "c": combined_hash, "a": accept_hashes,
                             "v": VERSIONS}, sort_keys=True)
build_id = sha(build_material)[:16]

payload = {
    "build_id": build_id,
    "algorithm": "sha256",
    "normalization": "authoritative .docx paragraphs+cells, stripped, joined by \\n",
    "document_versions": VERSIONS,
    "documents": doc_hashes,
    "combined_suite": combined_hash,
    "acceptance_text": accept_hashes,
}
json.dump(payload, open(OUT, "w", encoding="utf-8"), indent=2)
print("wrote", OUT, "build_id", build_id)
for k, h in doc_hashes.items():
    print(f"  {k}: {h[:16]}… v{VERSIONS[k]}")
print(f"  combined_suite: {combined_hash[:16]}…")
