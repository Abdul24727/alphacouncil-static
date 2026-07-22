"""Generate the English legal pages from the authoritative FINAL .docx + the central registry.
Faithful: text comes verbatim from the .docx (no rewriting). Emits directory-based clean routes.
"""
import os, re, html, importlib.util

DOCX_DIR = r"C:\projects\alphacouncil-static\legal\source"  # amended, version-controlled source
OUT_ROOT = r"C:\projects\alphacouncil-static\legal"
REG_PATH = r"C:\projects\alphacouncil\Argaam-council-before-testing-cluade-project-\llm-council\backend\legal_documents.py"

# Load the canonical registry standalone (no backend package import).
_spec = importlib.util.spec_from_file_location("legal_documents", REG_PATH)
REG = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(REG)

from docx import Document

CANON = "https://alphacouncil.tech"
ENTITY = "Establishment Insight Lane"
UNN = "7050149686"

CSS = """
:root{--bg:#07111F;--surface:#111A2B;--surface2:#141D2F;--fg:#F8FAFC;--fg-muted:#AAB6C5;
--fg-dim:#6B7A8F;--accent:#27D3C3;--border:rgba(255,255,255,.10);--maxw:820px}
*,*::before,*::after{box-sizing:border-box;margin:0;padding:0}
html{scroll-behavior:smooth}
body{font-family:'Inter',-apple-system,system-ui,sans-serif;background:var(--bg);color:var(--fg);line-height:1.65}
a{color:var(--accent);text-decoration:none}a:hover{text-decoration:underline}
.skip{position:absolute;left:-999px;top:0;background:var(--accent);color:#04121a;padding:10px 16px;border-radius:6px;z-index:100}
.skip:focus{left:12px;top:12px}
header.site{position:sticky;top:0;background:rgba(7,17,31,.92);backdrop-filter:blur(8px);border-bottom:1px solid var(--border);z-index:20}
.bar{max-width:var(--maxw);margin:0 auto;display:flex;align-items:center;gap:20px;padding:14px 20px;flex-wrap:wrap}
.logo{font-weight:700;font-size:18px;color:var(--fg)}
nav.legalnav{display:flex;gap:14px;flex-wrap:wrap;font-size:13px}
nav.legalnav a{color:var(--fg-muted)}nav.legalnav a[aria-current="page"]{color:var(--accent);font-weight:600}
main{max-width:var(--maxw);margin:0 auto;padding:32px 20px 64px}
.crumb{font-size:13px;color:var(--fg-dim);margin-bottom:14px}
h1{font-size:30px;line-height:1.25;margin-bottom:12px}
.docmeta{font-size:13px;color:var(--fg-muted);border:1px solid var(--border);background:var(--surface);border-radius:10px;padding:12px 16px;margin-bottom:8px;display:flex;flex-wrap:wrap;gap:6px 20px}
.docmeta b{color:var(--fg)}
.operator{font-size:13px;color:var(--fg-dim);margin:8px 0 26px}
.toc{border:1px solid var(--border);background:var(--surface);border-radius:10px;padding:16px 20px;margin-bottom:34px}
.toc h2{font-size:13px;text-transform:uppercase;letter-spacing:.06em;color:var(--fg-dim);margin-bottom:10px}
.toc ol{margin:0;padding-left:20px;columns:2;column-gap:28px}.toc li{margin:4px 0;font-size:14px}
@media(max-width:640px){.toc ol{columns:1}}
h2.sec{font-size:20px;margin:34px 0 12px;padding-top:8px;scroll-margin-top:80px}
h3.sub{font-size:16px;margin:22px 0 8px;color:var(--fg)}
p{margin:0 0 14px;color:var(--fg-muted)}p strong{color:var(--fg)}
p.def{margin:0 0 10px;color:var(--fg-muted)}
ul.doc,ol.doc{margin:0 0 16px;padding-left:26px}ul.doc li,ol.doc li{margin:6px 0;color:var(--fg-muted)}
table.doc{width:100%;border-collapse:collapse;margin:0 0 18px;font-size:14px}
table.doc th,table.doc td{border:1px solid var(--border);padding:8px 10px;text-align:left;vertical-align:top;color:var(--fg-muted)}
table.doc th{background:var(--surface2);color:var(--fg)}
.backtop{font-size:12px;margin-left:10px;color:var(--fg-dim)}
footer.site{border-top:1px solid var(--border);background:var(--surface);margin-top:40px}
.finner{max-width:var(--maxw);margin:0 auto;padding:26px 20px;font-size:13px;color:var(--fg-dim)}
.finner a{color:var(--fg-muted)}.flinks{display:flex;gap:16px;flex-wrap:wrap;margin-bottom:12px}
:focus-visible{outline:2px solid var(--accent);outline-offset:2px;border-radius:4px}
@media print{header.site,footer.site .flinks,.skip,.backtop{display:none}
 body{background:#fff;color:#000}main{max-width:100%}p,.docmeta,.operator,li,td,th{color:#000}
 h2.sec{page-break-after:avoid}.toc{break-inside:avoid}a{color:#000;text-decoration:underline}}
"""

def esc(s): return html.escape(s, quote=False)

from docx.text.run import Run as _Run
try:
    from docx.text.hyperlink import Hyperlink as _Hyperlink
except Exception:
    _Hyperlink = ()


def runs_html(p):
    # Iterate inner content so HYPERLINK text (e.g. mailto: emails) is included — p.runs omits it.
    # Emails/links render clickable; bold is preserved.
    try:
        items = list(p.iter_inner_content())
    except Exception:
        return esc(p.text)
    parts = []
    for item in items:
        t = esc(item.text)
        if not t:
            continue
        if _Hyperlink and isinstance(item, _Hyperlink):
            try:
                href = esc(item.address or "")
            except Exception:
                href = ""   # broken/internal rel — render the text, not a link
            parts.append(f'<a href="{href}">{t}</a>' if href else t)
        elif isinstance(item, _Run) and item.bold:
            parts.append(f"<strong>{t}</strong>")
        else:
            parts.append(t)
    joined = "".join(parts)
    return joined if joined.strip() else esc(p.text)

def render_body(doc):
    """Return (sections list [(id,title)], html_body) preserving structure."""
    parts, secs = [], []
    n = 0; list_open = None
    def close_list():
        nonlocal list_open
        if list_open: parts.append(f"</{list_open}>"); list_open = None
    for p in doc.paragraphs:
        st = (p.style.name if p.style else "") or ""
        txt = p.text.strip()
        if st == "Title":
            continue
        if st == "Legal Metadata":
            continue  # rendered from registry in the header
        if not txt:
            close_list(); continue
        if st.startswith("Heading 1"):
            close_list(); n += 1
            sid = f"sec-{n}"
            secs.append((sid, txt))
            parts.append(f'<h2 class="sec" id="{sid}">{runs_html(p)}</h2>')
        elif st.startswith("Heading 2") or st.startswith("Heading 3"):
            close_list(); parts.append(f'<h3 class="sub">{runs_html(p)}</h3>')
        elif st == "Definition":
            close_list(); parts.append(f'<p class="def">{runs_html(p)}</p>')
        elif "List Number" in st:
            if list_open != "ol": close_list(); parts.append('<ol class="doc">'); list_open = "ol"
            parts.append(f"<li>{runs_html(p)}</li>")
        elif "List" in st or "Bullet" in st:
            if list_open != "ul": close_list(); parts.append('<ul class="doc">'); list_open = "ul"
            parts.append(f"<li>{runs_html(p)}</li>")
        else:
            close_list(); parts.append(f"<p>{runs_html(p)}</p>")
    close_list()
    # tables (appended after prose; the FINAL docs use tables mainly in Privacy for data lists)
    for tbl in doc.tables:
        rows = tbl.rows
        if not rows: continue
        thtml = ['<table class="doc">']
        for ri, row in enumerate(rows):
            tag = "th" if ri == 0 else "td"
            thtml.append("<tr>" + "".join(f"<{tag}>{esc(c.text.strip())}</{tag}>" for c in row.cells) + "</tr>")
        thtml.append("</table>")
        parts.append("".join(thtml))
    return secs, "\n".join(parts)

NAV_ITEMS = [("terms","Terms"),("privacy","Privacy"),("refunds","Refunds"),
             ("cookies","Cookies"),("ai-financial-disclaimer","AI Disclaimer")]

def legalnav(active):
    a = []
    for k, label in NAV_ITEMS:
        cur = ' aria-current="page"' if k == active else ""
        a.append(f'<a href="{REG.LEGAL_DOCUMENTS[k]["route"]}"{cur}>{label}</a>')
    a.append(f'<a href="/legal"{" aria-current=\"page\"" if active=="index" else ""}>Legal index</a>')
    return "\n".join(a)

def footer():
    links = " ".join(f'<a href="{REG.LEGAL_DOCUMENTS[k]["route"]}">{lab}</a>' for k, lab in NAV_ITEMS)
    return f"""<footer class="site"><div class="finner">
<div class="flinks">{links} <a href="/legal">Legal index</a></div>
<p>{esc(ENTITY)} — a Saudi sole establishment, Unified National Number {UNN}, trading as AlphaCouncil.
Support: <a href="mailto:support@alphacouncil.tech">support@alphacouncil.tech</a> ·
Privacy: <a href="mailto:privacy@alphacouncil.tech">privacy@alphacouncil.tech</a></p>
<p>&copy; 2026 AlphaCouncil. All rights reserved.</p></div></footer>"""

def page(key, meta, secs, body):
    route = meta["route"]
    toc = ""
    if len(secs) > 4:
        items = "\n".join(f'<li><a href="#{sid}">{esc(t)}</a></li>' for sid, t in secs)
        toc = f'<nav class="toc" aria-label="Table of contents"><h2>On this page</h2><ol>{items}</ol></nav>'
    desc = f'{meta["title"]} for AlphaCouncil — research and analytics for GCC-listed equities. Informational analysis only; not licensed financial advice.'
    # Wrap the SUBSTANTIVE legal content in #legal-body with the canonical content hash embedded,
    # so the build check can prove the rendered body reproduces the authoritative source exactly.
    canonical = REG.content_hash(key) or (REG.combined_suite_hash() if key == "suite" else None)
    legal_body = (f'<div id="legal-body" data-canonical-hash="{canonical}">\n{body}\n</div>'
                  if canonical else body)
    return f"""<!DOCTYPE html>
<html lang="en" dir="ltr">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{esc(meta["title"])} — AlphaCouncil</title>
<meta name="description" content="{esc(desc)}">
<meta name="robots" content="index, follow">
<link rel="canonical" href="{CANON}{route}">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
<style>{CSS}</style>
</head>
<body id="top">
<a class="skip" href="#content">Skip to content</a>
<header class="site"><div class="bar">
<a class="logo" href="/index.html">AlphaCouncil</a>
<nav class="legalnav" aria-label="Legal documents">{legalnav(key)}</nav>
</div></header>
<main id="content">
<p class="crumb"><a href="/index.html">Home</a> › <a href="/legal">Legal</a> › {esc(meta["title"])}</p>
<h1>{esc(meta["title"])}</h1>
<div class="docmeta">
<span>{esc(REG.SUITE_LABEL)}</span>
<span><b>Version</b> {esc(meta["version"])}</span>
<span><b>Effective date</b> {esc(meta["effective_date"])}</span>
<span><b>Last updated</b> {esc(meta["effective_date"])}</span>
</div>
<p class="operator">Operator: {esc(ENTITY)}, a Saudi sole establishment, Unified National Number {UNN}, trading as AlphaCouncil.</p>
{toc}
{legal_body}
</main>
{footer()}
</body>
</html>"""

def write(path, content):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    open(path, "w", encoding="utf-8").write(content)

built = []
for key, meta in REG.LEGAL_DOCUMENTS.items():
    doc = Document(os.path.join(DOCX_DIR, meta["source"]))
    secs, body = render_body(doc)
    out = os.path.join(OUT_ROOT, key, "index.html")
    write(out, page(key, meta, secs, body))
    built.append((meta["route"], out, len(secs), len(body)))

# Combined suite page (/legal/suite) from the combined docx
cs = REG.COMBINED_SUITE
sdoc = Document(os.path.join(DOCX_DIR, cs["source"]))
ssecs, sbody = render_body(sdoc)
write(os.path.join(OUT_ROOT, "suite", "index.html"),
      page("suite", {**cs, "title": "AlphaCouncil Legal Suite", "route": "/legal/suite"}, ssecs, sbody))
built.append(("/legal/suite", "suite/index.html", len(ssecs), len(sbody)))

# /legal index
cards = []
for key, meta in REG.LEGAL_DOCUMENTS.items():
    cards.append(f'<li><a href="{meta["route"]}"><b>{esc(meta["title"])}</b></a> — Version {meta["version"]}, effective {meta["effective_date"]}</li>')
index_body = ('<p>These are the current legal documents that govern your use of AlphaCouncil. '
              'Each is available as its own page below; a combined document is also provided.</p>'
              '<ul class="doc">' + "\n".join(cards) + '</ul>'
              '<h2 class="sec" id="combined">Combined document</h2>'
              '<p><a href="/legal/suite">View the complete AlphaCouncil Legal Suite</a> — all documents in one page '
              '(same versions and effective dates as above).</p>')
write(os.path.join(OUT_ROOT, "index.html"),
      page("index", {"title": "Legal", "version": "1.0", "effective_date": "2026-07-21", "route": "/legal"},
           [("combined", "Combined document")], index_body))
built.append(("/legal", "index.html", 0, len(index_body)))

# Emit a public JSON registry for the frontend / checkout.
import json
write(os.path.join(OUT_ROOT, "registry.json"), json.dumps(REG.as_public_registry(), indent=2, ensure_ascii=False))

print("Generated pages:")
for route, out, nsec, blen in built:
    print(f"  {route:38s} sections={nsec:2d} bytes={blen}")
print("  /legal/registry.json")
